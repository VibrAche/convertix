"""Génère des fichiers d'exemple dans un dossier pour tester les conversions.

    python tests/make_samples.py <dossier>

Le DNG n'est pas généré (format RAW) : place un vrai DNG d'iPhone dans
`tests/samples_raw/` pour qu'il soit testé (voir CLAUDE.md).
"""
from __future__ import annotations

import subprocess
import sys
from pathlib import Path


def make(folder: Path) -> list[Path]:
    from PIL import Image, ImageDraw

    folder.mkdir(parents=True, exist_ok=True)
    made: list[Path] = []

    # Image RGBA avec transparence
    img = Image.new("RGBA", (640, 480), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    d.ellipse((80, 60, 560, 420), fill=(40, 120, 220, 255))
    d.text((260, 230), "Convertix", fill=(255, 255, 255, 255))
    img.save(folder / "transparent.png")
    made.append(folder / "transparent.png")

    # JPEG portrait avec orientation EXIF = 6 (comme un iPhone tenu verticalement)
    rgb = Image.new("RGB", (400, 300), (200, 80, 60))
    ImageDraw.Draw(rgb).rectangle((0, 0, 100, 50), fill=(255, 255, 0))
    exif = rgb.getexif()
    exif[0x0112] = 6
    exif[0x010F] = "Apple"
    rgb.save(folder / "orientation.jpg", exif=exif.tobytes(), quality=90)
    made.append(folder / "orientation.jpg")

    # HEIC (format photo iPhone)
    try:
        import pillow_heif
        pillow_heif.register_heif_opener()
        rgb.save(folder / "photo.heic", quality=80)
        made.append(folder / "photo.heic")
    except ImportError:
        pass

    # SVG (vectoriel, avec transparence)
    (folder / "logo.svg").write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" width="120" height="80" viewBox="0 0 120 80">'
        '<rect x="10" y="10" width="100" height="60" rx="14" fill="#0A84FF"/>'
        '<circle cx="60" cy="40" r="18" fill="white" fill-opacity="0.8"/></svg>', encoding="utf-8")
    made.append(folder / "logo.svg")

    # GIF animé
    frames = [Image.new("RGB", (120, 120), (i * 40 % 255, 100, 200)) for i in range(6)]
    frames[0].save(folder / "anim.gif", save_all=True, append_images=frames[1:], duration=100, loop=0)
    made.append(folder / "anim.gif")

    # PDF 2 pages
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz
    doc = fitz.open()
    for i in range(2):
        page = doc.new_page()
        page.insert_text((72, 100), f"Page {i + 1} — test Convertix éàç", fontsize=20)
    doc.save(str(folder / "document.pdf"))
    doc.close()
    made.append(folder / "document.pdf")

    # EPUB minimal (livre numérique)
    import zipfile
    with zipfile.ZipFile(folder / "livre.epub", "w") as z:
        z.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        z.writestr("META-INF/container.xml",
                   '<?xml version="1.0"?><container version="1.0" '
                   'xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles>'
                   '<rootfile full-path="content.opf" media-type="application/oebps-package+xml"/>'
                   '</rootfiles></container>')
        z.writestr("content.opf",
                   '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0" '
                   'unique-identifier="id"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/">'
                   '<dc:identifier id="id">convertix-test</dc:identifier><dc:title>Test</dc:title>'
                   '<dc:language>fr</dc:language></metadata><manifest>'
                   '<item id="c1" href="c1.xhtml" media-type="application/xhtml+xml"/></manifest>'
                   '<spine><itemref idref="c1"/></spine></package>')
        z.writestr("c1.xhtml",
                   '<?xml version="1.0" encoding="utf-8"?><html xmlns="http://www.w3.org/1999/xhtml">'
                   '<head><title>Chapitre 1</title></head><body><h1>Chapitre 1</h1>'
                   '<p>Il était une fois un convertisseur très pratique. éàç</p></body></html>')
    made.append(folder / "livre.epub")

    # Word
    import docx
    dd = docx.Document()
    dd.add_heading("Test Convertix", 1)
    dd.add_paragraph("Bonjour, ceci est un document Word avec des accents : éèàùç.")
    t = dd.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "A", "B", "1", "2"
    dd.save(folder / "lettre.docx")
    made.append(folder / "lettre.docx")

    # Excel (2 feuilles) + CSV à la française
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Ventes"
    ws.append(["Mois", "Montant"])
    ws.append(["Janvier", 1200.5])
    ws2 = wb.create_sheet("Clients")
    ws2.append(["Nom", "Ville"])
    ws2.append(["Dupont", "Grasse"])
    wb.save(folder / "tableau.xlsx")
    made.append(folder / "tableau.xlsx")
    (folder / "export.csv").write_text("Nom;Montant\nÉlodie;12,50\nMarc;7\n", encoding="cp1252")
    made.append(folder / "export.csv")

    # Vidéo + audio via ffmpeg
    try:
        import imageio_ffmpeg
        ff = imageio_ffmpeg.get_ffmpeg_exe()
        subprocess.run([ff, "-y", "-loglevel", "error", "-f", "lavfi", "-i", "testsrc=duration=2:size=320x240:rate=15",
                        "-f", "lavfi", "-i", "sine=frequency=440:duration=2", "-shortest",
                        "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac", str(folder / "video.mov")], check=True)
        made.append(folder / "video.mov")
        subprocess.run([ff, "-y", "-loglevel", "error", "-f", "lavfi", "-i", "sine=frequency=440:duration=2",
                        str(folder / "son.wav")], check=True)
        made.append(folder / "son.wav")
    except Exception as e:  # pragma: no cover
        print("ffmpeg indisponible :", e)

    return made


if __name__ == "__main__":
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "samples")
    for p in make(out):
        print(p)
